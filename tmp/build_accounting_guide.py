from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/docs/SVdP Voucher System How-To Guide - Accounting and Bookkeeping.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 88, 88)
INK = RGBColor(0, 0, 0)
LIGHT_BLUE = "E8EEF5"
PALE = "F4F6F9"
WARNING = "FFF2CC"


def set_run_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para_format(p, before=0, after=6, line=1.2, left=None, first=None, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if left is not None:
        p.paragraph_format.left_indent = Inches(left)
    if first is not None:
        p.paragraph_format.first_line_indent = Inches(first)
    p.paragraph_format.keep_together = keep
    p.paragraph_format.keep_with_next = keep


def add_para(doc, text="", size=10.5, bold=False, color=INK, before=0, after=6, align=None):
    p = doc.add_paragraph()
    para_format(p, before, after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size, bold, color)
    return p


def add_title_block(doc):
    add_para(doc, "FOR ACCOUNTING & BOOKKEEPING USE", 9, True, MUTED, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "THE SOCIETY OF ST. VINCENT DE PAUL", 14, True, DARK_BLUE, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Neighbors helping Neighbors", 10.5, False, MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Voucher System", 24, True, INK, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "How-To Guide for Accounting & Bookkeeping", 15, False, BLUE, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 7
    elif level == 2:
        size, color, before, after = 13, BLUE, 10, 5
    else:
        size, color, before, after = 11.5, DARK_BLUE, 8, 3
    para_format(p, before, after, line=1.15, keep=True)
    r = p.add_run(text)
    set_run_font(r, size, True, color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    indent = 0.44 + (level * 0.24)
    para_format(p, 0, 3, line=1.18, left=indent, first=-0.22)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_step(doc, number, text):
    p = doc.add_paragraph()
    para_format(p, 0, 3, line=1.18, left=0.05)
    label = p.add_run(f"Step {number}: ")
    set_run_font(label, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=9.45, bold=False, color=INK):
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.12)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for cell, head in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        set_cell_text(cell, head, 9.4, True, DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_data):
            set_cell_text(cell, text)
    return table


def add_note(doc, label, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.18)
    r = p.add_run(label + ": ")
    set_run_font(r, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_quick_reference(doc):
    doc.add_page_break()
    add_heading(doc, "Quick Reference Card", 1)
    add_para(doc, "Use this as a month-end reminder after you have read the full guide.", 10.5, False, MUTED, after=8)
    rows = [
        ("Analytics review", "Analytics tab -> set filters -> Apply Filters -> review totals, organization results, denied vouchers, and overrides."),
        ("Export voucher data", "Analytics tab -> Export Data -> choose range -> choose whether to include denied vouchers -> Export to Excel."),
        ("Find invoices", "Invoices tab -> choose conference, dates, and statement status -> Apply Filters."),
        ("Open invoice", "Invoices tab -> Invoice Results -> Open Invoice."),
        ("Preview statement", "Statements tab -> choose conference and period -> Preview Eligible Invoices."),
        ("Generate statement", "Statements tab -> review preview -> Generate Statement -> confirm -> Open Statement."),
        ("Update bookkeeping settings", "Accounting tab -> enter bookkeeping email and QuickBooks names -> Save Accounting Settings."),
        ("Run month-end cycle", "Accounting tab -> Run Due Monthly Cycle Now -> review Recent Accounting Batches."),
        ("Download files", "Accounting tab -> Recent Accounting Batches -> IIF or Manifest."),
    ]
    add_table(doc, ["TASK", "WHERE / WHAT TO DO"], rows, [2450, 6910])
    add_heading(doc, "Important Rules", 2)
    add_bullet(doc, "Invoices are created when furniture or household goods fulfillment is completed.")
    add_bullet(doc, "Statements only include invoices that have not already been placed on a statement.")
    add_bullet(doc, "The usual statement dates default to the first and last day of the previous month.")
    add_bullet(doc, "The Accounting cycle creates statements, sends statement emails when possible, and creates the QuickBooks package.")
    add_bullet(doc, "Use the Manifest to check what was included before or after importing the IIF file into QuickBooks.")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.45)
section.footer_distance = Inches(0.38)

for style_name in ["Normal", "List Bullet", "List Number"]:
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_run = header.add_run("SVdP Voucher System - Accounting & Bookkeeping Guide")
set_run_font(header_run, 8.5, False, MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("(c) 2026 Society of St. Vincent de Paul - Fort Wayne. All rights reserved.")
set_run_font(footer_run, 8.5, False, MUTED)

add_title_block(doc)
add_heading(doc, "What This Guide Covers", 1)
add_para(doc, "This guide explains the Voucher System tabs used for accounting and bookkeeping work.", 10.5)
add_bullet(doc, "Analytics is used to review voucher activity, totals, trends, denied vouchers, overrides, and exports.")
add_bullet(doc, "Invoices is used to find stored invoices and see whether they have already been placed on a statement.")
add_bullet(doc, "Statements is used to create stored statements for a Conference or Partner and date range.")
add_bullet(doc, "Accounting is used for bookkeeping settings, organization billing mappings, monthly batches, and QuickBooks files.")
add_note(doc, "Important", "This guide does not cover catalog management, cashier redemption steps, or manager approval codes.")

add_heading(doc, "Accounting Tabs at a Glance", 1)
add_table(doc, ["TAB", "WHAT IT IS FOR"], [
    ("Analytics", "Voucher counts, redemption totals, people served, coats issued, denied vouchers, override statistics, and Excel export."),
    ("Invoices", "Stored invoice lookup by Conference, invoice date, and statement status."),
    ("Statements", "Statement creation for one Conference and period, eligible invoice preview, and recent statement lookup."),
    ("Accounting", "Bookkeeping email, QuickBooks mapping, organization billing details, monthly cycle, and batch files."),
], [2300, 7060])

add_heading(doc, "Before Month-End", 1)
add_bullet(doc, "Make sure each organization that should receive statements has a Billing Email when possible.")
add_bullet(doc, "Make sure each organization has the correct QuickBooks Customer Name if it will be included in QuickBooks exports.")
add_bullet(doc, "Make sure the Accounting Settings have the correct QuickBooks account and item names.")
add_bullet(doc, "Review unstatemented invoices before generating or running the monthly cycle.")

doc.add_page_break()
add_heading(doc, "Analytics Tab", 1)
add_para(doc, "The Analytics tab is the reporting view. It is used to understand activity and prepare exports.", 10.5)
add_heading(doc, "Use the Filters", 2)
add_steps(doc, [
    "Open the Analytics tab.",
    "Choose the Date Range: All Time, Last 30 Days, Last 90 Days, Last Year, Year to Date, or Custom Range.",
    "If you choose Custom Range, enter the start date and end date.",
    "Choose Organization Type if needed: Conference, Partner, Store, or All Types.",
    "Choose a Specific Organization if needed.",
    "Choose Voucher Type if needed.",
    "Click Apply Filters.",
])
add_bullet(doc, "The selected filters update the metrics on the page.")
add_bullet(doc, "Use Reset to return to the default view.")

add_heading(doc, "What the Analytics Sections Show", 2)
add_table(doc, ["SECTION", "WHAT IT TELLS YOU"], [
    ("Overview Statistics", "Total vouchers, active vouchers, redeemed vouchers, and denied or blocked vouchers."),
    ("Time Period Analysis", "Voucher activity for the last 30 days, last 90 days, and current year."),
    ("Community Impact", "People served, total value provided, and coat counts."),
    ("Items Provided", "Items redeemed, redemption value, and average items per voucher."),
    ("Organization Type Breakdown", "Voucher and redemption activity by organization type."),
    ("Breakdown by Voucher Type", "Created vouchers, redeemed items, and value by voucher type."),
    ("Performance by Organization", "Redemption results for each Conference, Partner, and Store."),
    ("Denied/Blocked Vouchers", "Blocked requests, recent denied vouchers, and denial reasons."),
    ("Emergency Override Statistics", "Total overrides, overrides by manager, and overrides by reason."),
], [2850, 6510])

add_heading(doc, "Export Data", 2)
add_steps(doc, [
    "Open the Analytics tab.",
    "Scroll to Export Data.",
    "Choose the Date Range for the export.",
    "If you choose Custom Range, enter the start date and end date.",
    "Check Include blocked/denied vouchers in export only if those records should be included.",
    "Click Export to Excel.",
])
add_bullet(doc, "The export includes voucher details, Conference or Partner, organization type, voucher type, status, redeemed date, redeemed items, redemption value, coat details, override manager, and override reason.")
add_bullet(doc, "Use this export for reporting and review. Use the Accounting tab for QuickBooks package files.")

doc.add_page_break()
add_heading(doc, "Invoices Tab", 1)
add_para(doc, "The Invoices tab lists stored invoices. Invoices are created when furniture or household goods fulfillment is completed.", 10.5)
add_note(doc, "What an invoice includes", "An invoice shows the Conference, invoice date, neighbor, voucher number, fulfilled items, item total, Conference share, delivery fee when charged, and total invoice amount.", PALE)
add_heading(doc, "Filter Stored Invoices", 2)
add_steps(doc, [
    "Open the Invoices tab.",
    "Choose a Conference, or leave it set to All conferences.",
    "Enter Invoice Date From if you want a start date.",
    "Enter Invoice Date To if you want an end date.",
    "Choose Statement Status: All invoices, Unstatemented only, or Already statemented.",
    "Click Apply Filters.",
])
add_bullet(doc, "Use Reset to clear the filters and reload the invoice list.")
add_bullet(doc, "The summary shows how many invoices were found and the total amount.")

add_heading(doc, "Read the Invoice Results", 2)
add_table(doc, ["COLUMN", "WHAT IT MEANS"], [
    ("Invoice", "The stored invoice number."),
    ("Conference", "The Conference or Partner connected to the invoice."),
    ("Neighbor", "The neighbor connected to the voucher."),
    ("Date", "The invoice date."),
    ("Amount", "The total amount billed on that invoice."),
    ("Statement", "The statement number, or Unstatemented if it has not been placed on a statement yet."),
    ("Document", "Open Invoice, or Missing file if the stored document cannot be found."),
], [2250, 7110])

add_heading(doc, "Open an Invoice", 2)
add_steps(doc, [
    "Find the invoice in Invoice Results.",
    "Click Open Invoice.",
    "Review the invoice details in the document that opens.",
])
add_bullet(doc, "Use Unstatemented only when preparing invoices for a new statement.")
add_bullet(doc, "Use Already statemented when checking whether a prior invoice was already included.")

doc.add_page_break()
add_heading(doc, "Statements Tab", 1)
add_para(doc, "The Statements tab creates stored statements for one Conference or Partner and one date range.", 10.5)
add_note(doc, "Important", "Only invoices that do not already belong to a statement are eligible. This helps prevent the same invoice from being billed twice.", WARNING)

add_heading(doc, "Preview Eligible Invoices", 2)
add_steps(doc, [
    "Open the Statements tab.",
    "Choose the Conference.",
    "Review the Period Start and Period End dates.",
    "Click Preview Eligible Invoices.",
])
add_bullet(doc, "The system shows the eligible invoice count and total amount.")
add_bullet(doc, "The preview list shows invoice number, neighbor, date, amount, and document link.")
add_bullet(doc, "If no invoices appear, there are no unstatemented invoices for that Conference and period.")

add_heading(doc, "Generate a Statement", 2)
add_steps(doc, [
    "Preview the eligible invoices first.",
    "Confirm the Conference and date range are correct.",
    "Click Generate Statement.",
    "Confirm when the system asks if you want to generate the statement.",
    "Wait for the success message.",
    "Click Open Statement if you want to review the statement immediately.",
])
add_bullet(doc, "A generated statement receives a statement number.")
add_bullet(doc, "The invoices included on that statement are marked as statemented.")
add_bullet(doc, "After the statement is generated, those same invoices will no longer appear in an Unstatemented only search.")

add_heading(doc, "Recent Statements", 2)
add_table(doc, ["COLUMN", "WHAT IT MEANS"], [
    ("Statement", "The stored statement number."),
    ("Conference", "The Conference or Partner on the statement."),
    ("Period", "The date range covered by the statement."),
    ("Invoices", "How many invoices are included."),
    ("Total", "The statement total."),
    ("Generated", "When the statement was created."),
    ("Document", "Open Statement, or Missing file if the stored document cannot be found."),
], [2250, 7110])

doc.add_page_break()
add_heading(doc, "Accounting Tab", 1)
add_para(doc, "The Accounting tab controls bookkeeping settings and monthly accounting batches.", 10.5)
add_heading(doc, "Accounting Settings", 2)
add_steps(doc, [
    "Open the Accounting tab.",
    "Enter the Bookkeeping Email.",
    "Enter the QuickBooks Accounts Receivable Account.",
    "Enter the QuickBooks Income Account.",
    "Enter the QuickBooks Voucher Service Item.",
    "Enter the QuickBooks Delivery Service Item.",
    "Click Save Accounting Settings.",
])
add_bullet(doc, "The Bookkeeping Email receives the QuickBooks package when the system emails a batch.")
add_bullet(doc, "QuickBooks names should match the names used in QuickBooks.")
add_note(doc, "Use care", "If the QuickBooks account or item names are wrong, the import file may not land where bookkeeping expects it.", WARNING)

add_heading(doc, "Organization Billing Mappings", 2)
add_steps(doc, [
    "Open the Accounting tab.",
    "Find the organization you need to update.",
    "Enter or update the Billing Email.",
    "Enter or update the QuickBooks Customer Name.",
    "Click Save on that organization row.",
])
add_bullet(doc, "The Billing Email is used for statement email delivery when available.")
add_bullet(doc, "The QuickBooks Customer Name is used in the QuickBooks import file.")
add_bullet(doc, "If a Billing Email is missing, the system may fall back to the organization's notification email when sending a statement.")

add_heading(doc, "Run Due Monthly Cycle Now", 2)
add_para(doc, "The monthly cycle is the bookkeeping close-out process. It gathers unstatemented invoices through the end of the previous month.", 10.5)
add_steps(doc, [
    "Open the Accounting tab.",
    "Confirm Accounting Settings and Organization Billing Mappings are ready.",
    "Click Run Due Monthly Cycle Now.",
    "Wait for the page to return to the Accounting tab.",
    "Review Recent Accounting Batches.",
])
add_bullet(doc, "The monthly cycle creates statements by organization when eligible invoices exist.")
add_bullet(doc, "It creates statement PDFs when the PDF tool is available.")
add_bullet(doc, "It sends statements when an email address is available and mail delivery works.")
add_bullet(doc, "It creates the QuickBooks package and manifest for the batch.")

doc.add_page_break()
add_heading(doc, "Recent Accounting Batches", 1)
add_para(doc, "Recent Accounting Batches show the results of monthly accounting runs.", 10.5)
add_table(doc, ["COLUMN", "WHAT IT MEANS"], [
    ("Batch", "The batch name, usually tied to the month being closed."),
    ("Cutoff", "The final invoice date included in the batch."),
    ("Status", "Processing, completed, partial, or blocked."),
    ("Statements", "How many statements were included."),
    ("Invoices", "How many invoices were included."),
    ("Total", "The total dollar amount included in the batch."),
    ("Email", "Whether the bookkeeping package email was sent, failed, or has not been sent."),
    ("Actions", "Download IIF, download Manifest, or Re-email when files are available."),
], [1900, 7460])

add_heading(doc, "Download the QuickBooks File", 2)
add_steps(doc, [
    "Open the Accounting tab.",
    "Find the batch in Recent Accounting Batches.",
    "Click IIF.",
    "Use the downloaded IIF file for the QuickBooks import process.",
])
add_bullet(doc, "The IIF file contains invoice lines for QuickBooks.")
add_bullet(doc, "Delivery appears as its own line when a delivery fee was charged.")

add_heading(doc, "Download the Manifest", 2)
add_steps(doc, [
    "Open the Accounting tab.",
    "Find the batch in Recent Accounting Batches.",
    "Click Manifest.",
    "Open the downloaded file to review what was included.",
])
add_bullet(doc, "The Manifest lists statement, invoice, organization, neighbor, amount, email status, and warnings.")
add_bullet(doc, "Use the Manifest as the plain-language check sheet for the batch.")

add_heading(doc, "Re-email a Batch", 2)
add_steps(doc, [
    "Open the Accounting tab.",
    "Find the batch in Recent Accounting Batches.",
    "Click Re-email.",
])
add_bullet(doc, "Use Re-email when the bookkeeping package needs to be sent again.")
add_bullet(doc, "If email continues to fail, download the IIF and Manifest directly from the batch row.")

add_heading(doc, "What the Batch Email Contains", 2)
add_bullet(doc, "One email includes the QuickBooks invoice import file and the statement manifest.")
add_bullet(doc, "Additional emails may include statement PDFs when there are too many or they are too large for one message.")

doc.add_page_break()
add_heading(doc, "Common Questions", 1)
qa = [
    ("Why is an invoice marked Unstatemented?", "It has not yet been placed on a statement."),
    ("Why did an invoice disappear from the statement preview?", "It was likely already added to a statement, or it no longer matches the selected Conference or date range."),
    ("Can one invoice be added to two statements?", "No. Once an invoice is attached to a statement, it is no longer eligible for a new one."),
    ("What should I do if a document says Missing file?", "The record exists, but the stored document cannot be opened from that row. Check with the system administrator before recreating bookkeeping work."),
    ("What does partial mean on a batch?", "Some work finished, but the system also recorded an issue, often with document creation or email delivery."),
    ("What does blocked mean on a batch?", "The batch could not finish, often because required accounting settings were missing."),
    ("Which file goes into QuickBooks?", "Use the IIF file for QuickBooks. Use the Manifest to review what was included."),
    ("Should I use Analytics export or Accounting export?", "Use Analytics export for reporting. Use the Accounting batch files for QuickBooks bookkeeping."),
]
for question, answer in qa:
    add_para(doc, question, 10.3, True, DARK_BLUE, before=4, after=2)
    add_para(doc, answer, 10.2, False, INK, after=4)

add_heading(doc, "Month-End Checklist", 1)
add_bullet(doc, "Review Analytics for the month and export voucher data if needed.")
add_bullet(doc, "Filter Invoices for Unstatemented only and check the expected date range.")
add_bullet(doc, "Preview Statements before generating them manually.")
add_bullet(doc, "Confirm Accounting Settings and Organization Billing Mappings.")
add_bullet(doc, "Run the monthly cycle or confirm the monthly batch has completed.")
add_bullet(doc, "Download the IIF and Manifest.")
add_bullet(doc, "Check the Manifest for warnings before finishing the bookkeeping process.")

add_quick_reference(doc)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
