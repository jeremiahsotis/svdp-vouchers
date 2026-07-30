from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/docs/SVdP Voucher System How-To Guide - Cashiers.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 88, 88)
INK = RGBColor(0, 0, 0)
LIGHT_BLUE = "E8EEF5"
PALE = "F4F6F9"
WARNING = "FFF2CC"


def set_run_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para_format(p, before=0, after=6, line=1.2, left=None, first=None, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if left is not None:
        p.paragraph_format.left_indent = Inches(left)
    if first is not None:
        p.paragraph_format.first_line_indent = Inches(first)
    p.paragraph_format.keep_together = keep


def add_para(doc, text="", size=10.5, bold=False, color=INK, before=0, after=6, align=None):
    p = doc.add_paragraph()
    para_format(p, before, after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size, bold, color)
    return p


def add_title_block(doc):
    add_para(doc, "FOR INTERNAL THRIFT STORE USE ONLY", 9, True, MUTED, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "THE SOCIETY OF ST. VINCENT DE PAUL", 14, True, DARK_BLUE, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Neighbors helping Neighbors", 10.5, False, MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Voucher System", 24, True, INK, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "How-To Guide for Cashiers", 15, False, BLUE, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 7
    elif level == 2:
        size, color, before, after = 13, BLUE, 10, 5
    else:
        size, color, before, after = 11.5, DARK_BLUE, 8, 3
    para_format(p, before, after, line=1.15, keep=True)
    r = p.add_run(text)
    set_run_font(r, size, True, color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    indent = 0.44 + (level * 0.24)
    para_format(p, 0, 3, line=1.18, left=indent, first=-0.22)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_step(doc, number, text):
    p = doc.add_paragraph()
    para_format(p, 0, 3, line=1.18, left=0.05)
    label = p.add_run(f"Step {number}: ")
    set_run_font(label, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=9.8, bold=False, color=INK):
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.12)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for cell, head in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        set_cell_text(cell, head, 9.5, True, DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_data):
            set_cell_text(cell, text, 9.5, False, INK)
    return table


def add_note(doc, label, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.18)
    r = p.add_run(label + ": ")
    set_run_font(r, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_quick_reference(doc):
    doc.add_page_break()
    add_heading(doc, "Quick Reference Card", 1)
    add_para(doc, "Use this as a station-side reminder after you have read the full guide.", 10.5, False, MUTED, after=8)
    rows = [
        ("Find a voucher", "Search by name, date of birth, or organization. Use filters for Active, Redeemed, Expired, or Coat Available."),
        ("Open details", "Click the voucher card. The full details and available actions appear on the right."),
        ("Redeem clothing", "Click Enter Items, enter adult and child item counts, then click Mark as Redeemed."),
        ("Issue coats", "Use only when the coat area says available. Click Issue Coat, enter actual coat counts, then click Issue Coats."),
        ("Emergency voucher", "Click Emergency Voucher, enter the household information, then click Create Emergency Voucher."),
        ("Furniture / household goods", "Resolve requested items or units, save progress as needed, then finalize when everything is resolved."),
    ]
    add_table(doc, ["ACTION", "WHAT TO DO"], rows, [2300, 7060])
    add_heading(doc, "Key Rules", 2)
    add_bullet(doc, "Vouchers expire 30 days after they are created.")
    add_bullet(doc, "Clothing vouchers are redeemed in one visit. Remaining items are not saved for another day.")
    add_bullet(doc, "Emergency clothing vouchers allow 3 clothing items per person.")
    add_bullet(doc, "Conference and Partner clothing vouchers allow 7 clothing items per person.")
    add_bullet(doc, "Coat eligibility resets every August 1st.")
    add_bullet(doc, "Do not record more items or coats than the neighbor actually receives.")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.45)
section.footer_distance = Inches(0.38)

for style_name in ["Normal", "List Bullet", "List Number"]:
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.text = "SVdP Voucher System - Cashier Guide"
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(header.runs[0], 8.5, False, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("(c) 2026 Society of St. Vincent de Paul - Fort Wayne. All rights reserved.")
set_run_font(r, 8.5, False, MUTED)

add_title_block(doc)

add_heading(doc, "What You Need to Know", 1)
add_para(doc, "The Cashier Station helps you find vouchers, check eligibility, redeem clothing vouchers, create emergency clothing vouchers, issue winter coats, and work through furniture or household goods requests when those appear.", 10.5)
add_bullet(doc, "You will see vouchers created by Conferences, Community Partners, and the store team.")
add_bullet(doc, "Most cashier work happens in two places: the voucher list on the left and the voucher details on the right.")
add_bullet(doc, "The screen refreshes automatically about every 30 seconds, so new vouchers appear without reloading the page.")
add_note(doc, "Important", "This guide covers ordinary cashier actions only. If the system does not give you an action, ask your supervisor what to do next.", WARNING)

add_heading(doc, "Opening the Cashier Station", 1)
add_steps(doc, [
    "Open Microsoft Edge or Google Chrome.",
    "Go to the Cashier Station page used by the store.",
    "Sign in if the page asks you to sign in again.",
    "Wait for the voucher list to load.",
])
add_para(doc, "If the screen says Re-authentication Required, click Sign In Again and return to the Cashier Station.", 10.3, False, MUTED)

add_heading(doc, "How the Screen Is Organized", 1)
add_table(doc, ["AREA", "WHAT IT DOES"], [
    ("Search box", "Finds vouchers by name, date of birth, or Conference / Partner organization."),
    ("Filter menu", "Shows All, Active, Redeemed, Expired, or Coat Available vouchers."),
    ("Sort menu", "Changes the order: newest, oldest, name A-Z, or name Z-A."),
    ("Voucher cards", "Shows the neighbor, organization, status, voucher type, household, and coat or fulfillment summary."),
    ("Detail panel", "Shows the full voucher and the actions available for that voucher."),
    ("Emergency Voucher button", "Opens the form to create a same-day emergency clothing voucher."),
], [2100, 7260])

add_heading(doc, "How to Look Up a Voucher", 1)
add_steps(doc, [
    "Click in the search box that says Search by name, DOB, or conference.",
    "Type the neighbor's last name, first name, date of birth, or organization name.",
    "Watch the list narrow down automatically.",
    "Use the filter menu if you only want Active, Redeemed, Expired, or Coat Available vouchers.",
    "Use the sort menu if you need oldest first, newest first, or names in order.",
    "Click the right voucher card to open the full details.",
])
add_heading(doc, "Search Tips", 2)
add_bullet(doc, "You do not need the whole name. Part of a name usually works.")
add_bullet(doc, "Check the date of birth before redeeming. Many neighbors can have similar names.")
add_bullet(doc, "If no vouchers match, clear the search box and try a simpler spelling.")
add_bullet(doc, "If a voucher is expired, you can view it, but you cannot use the ordinary redemption workflow.")

add_heading(doc, "Reading a Voucher Card", 1)
add_table(doc, ["FIELD", "WHAT IT MEANS"], [
    ("Status", "Ready, Redeemed, Expired, or a fulfillment status for furniture / household goods."),
    ("Neighbor", "The person receiving help."),
    ("Conference / Partner", "The organization that requested the voucher."),
    ("Voucher Type", "Clothing, Furniture, or Household Goods."),
    ("DOB", "The neighbor's date of birth. Use this to confirm the right person."),
    ("Household", "How many people are included in the request."),
    ("Coat", "For clothing vouchers only: Available, Issued, or Not Eligible."),
    ("Delivery", "For furniture or household goods: Yes or No."),
], [2100, 7260])

add_heading(doc, "Printing a Neighbor Voucher", 1)
add_steps(doc, [
    "Find and open the voucher.",
    "Click Print Neighbor Voucher in the detail panel.",
    "Print the voucher if the neighbor or store team needs a paper copy.",
])
add_para(doc, "A printed copy is helpful, but the voucher record in the system is still the source of truth.", 10.3, False, MUTED)

add_heading(doc, "Redeeming a Clothing Voucher", 1)
add_para(doc, "Use this when the neighbor has finished shopping for clothing and is ready to check out.", 10.5)
add_steps(doc, [
    "Find the neighbor in the voucher list.",
    "Click the voucher card to open the details.",
    "Confirm the name, date of birth, household size, status, and items allowed.",
    "Click Enter Items in the Redeem Voucher section.",
    "Enter Adult Items Provided and Child Items Provided.",
    "Check the Current total and Maximum total items shown on the screen.",
    "Click Mark as Redeemed.",
    "Wait for the success message and confirm the voucher status changes to Redeemed.",
])
add_heading(doc, "Clothing Item Rules", 2)
add_bullet(doc, "Conference and Partner clothing vouchers allow 7 items per person.")
add_bullet(doc, "Emergency clothing vouchers allow 3 items per person.")
add_bullet(doc, "Enter the actual number of items taken today.")
add_bullet(doc, "Do not enter more than the maximum shown by the system.")
add_bullet(doc, "If the neighbor takes fewer items than allowed, that is okay. Remaining items are not saved for another day.")

add_heading(doc, "Important Time Rules", 1)
add_heading(doc, "30 Days to Use", 2)
add_bullet(doc, "Vouchers expire 30 days after they are created.")
add_bullet(doc, "Expired vouchers are read-only in the ordinary cashier workflow.")
add_bullet(doc, "The list and detail panel show when a voucher is expired.")
add_heading(doc, "One Visit for Clothing", 2)
add_bullet(doc, "When a neighbor redeems a clothing voucher, complete the checkout that day.")
add_bullet(doc, "Do not hold part of the voucher for a later visit.")
add_bullet(doc, "Mark the voucher as Redeemed when checkout is complete.")
add_heading(doc, "90 Days Between Clothing Vouchers", 2)
add_bullet(doc, "The system checks recent vouchers when an emergency clothing voucher is created.")
add_bullet(doc, "If the system says the neighbor is not eligible, do not force the request through ordinary cashier steps.")

add_heading(doc, "Winter Coat Rules", 1)
add_para(doc, "Coats are separate from clothing item redemption. A neighbor can redeem clothing without receiving coats.", 10.5)
add_heading(doc, "Simple Rule", 2)
add_bullet(doc, "Each person in the household can receive one winter coat per coat year.")
add_bullet(doc, "The coat year resets every August 1st.")
add_bullet(doc, "The system shows whether coats are Available, Issued, or Not Eligible.")
add_heading(doc, "How to Issue Coats", 2)
add_steps(doc, [
    "Find and open the neighbor's clothing voucher.",
    "Look at the coat message in the voucher details.",
    "If the coat area says available, click Issue Coat.",
    "Enter the number of adult coats and children's coats actually given today.",
    "Check the total coat count shown on the screen.",
    "Click Issue Coats.",
    "Confirm the coat area now shows coats were issued.",
])
add_note(doc, "Remember", "Do not automatically issue coats for everyone in the household. Enter only the coats actually taken today.", PALE)

add_heading(doc, "Creating an Emergency Clothing Voucher", 1)
add_para(doc, "Use this when someone walks in needing clothing help right away and there is no Conference or Partner voucher ready for them.", 10.5)
add_steps(doc, [
    "Click Emergency Voucher at the top right of the Cashier Station.",
    "Enter the neighbor's First Name and Last Name.",
    "Enter the neighbor's Date of Birth.",
    "Enter the number of Adults and Children in the household.",
    "Click Create Emergency Voucher.",
    "Wait for the success message.",
    "The new emergency voucher will appear in the list and can be redeemed like a clothing voucher.",
])
add_heading(doc, "Emergency Voucher Amount", 2)
add_bullet(doc, "Emergency clothing vouchers allow 3 items per person.")
add_bullet(doc, "Example: 1 adult = 3 items.")
add_bullet(doc, "Example: 2 adults and 1 child = 9 items total.")
add_heading(doc, "If the System Says They Are Not Eligible", 2)
add_bullet(doc, "The neighbor may have received a recent voucher or may already have an active voucher.")
add_bullet(doc, "Do not create a second emergency voucher through ordinary cashier steps.")
add_bullet(doc, "Ask your supervisor what to do next.")

add_heading(doc, "Furniture Vouchers", 1)
add_para(doc, "Furniture vouchers are handled item by item. The screen shows requested items, delivery information, pricing fields, and whether each item is still unresolved.", 10.5)
add_heading(doc, "How to Review a Furniture Voucher", 2)
add_steps(doc, [
    "Search for the neighbor and open the furniture voucher.",
    "Check the voucher type badge, status, household, delivery, and requested item count.",
    "Review the Delivery Details section if delivery is requested.",
    "Use the Item Resolution section to work through each requested item.",
])
add_heading(doc, "Resolving Furniture Items", 2)
add_bullet(doc, "If the requested item is available, use Fulfill Item, enter the actual price, add notes if needed, and mark it completed.")
add_bullet(doc, "Photos are optional records. Add a photo when it helps document the item.")
add_bullet(doc, "If a different item is used instead, use Substitute and save the substitute item.")
add_bullet(doc, "If the item cannot be provided, use Cancel Item and choose the reason.")
add_bullet(doc, "When all items are resolved, use the completion action shown by the system to generate the final documents.")

add_heading(doc, "Household Goods Vouchers", 1)
add_para(doc, "Household Goods vouchers use a fulfillment workspace. The goal is to resolve every requested unit as fulfilled or unavailable.", 10.5)
add_steps(doc, [
    "Search for the neighbor and open the Household Goods voucher.",
    "Review the requested units, delivery information, and redeem-by date.",
    "For each line, enter price and quantity for fulfilled items.",
    "If some items are unavailable, enter the unavailable quantity and reason.",
    "Use Save Progress if the voucher is not finished yet.",
    "When every requested unit is resolved, click Finalize Voucher.",
])
add_note(doc, "For Furniture and Household Goods", "Each voucher remains independently redeemable and expirable, even when it is part of a larger request group.", PALE)

add_heading(doc, "Common Questions", 1)
qa = [
    ("What if I cannot find the voucher?", "Clear the search box and try the last name, first name, date of birth, or organization name. Also check the filter menu."),
    ("What if the voucher is expired?", "You can view it, but ordinary redemption is not available."),
    ("What if the neighbor lost the paper voucher?", "Search for the voucher in the system. If it is active and matches the neighbor, the system record can still be used."),
    ("Do coats have to be issued during clothing redemption?", "No. Coats are separate. Issue coats only when the neighbor is eligible and actually receives coats."),
    ("What if the neighbor only takes part of the allowed clothing items?", "Enter the actual item count and redeem the voucher. Remaining clothing items are not saved for later."),
    ("What if the screen stops updating?", "Wait a moment. If the session message says to sign in again, sign in again. If the problem continues, ask your supervisor."),
]
for q, a in qa:
    p = doc.add_paragraph()
    para_format(p, before=2, after=2, line=1.16)
    r = p.add_run("Q: " + q + " ")
    set_run_font(r, 10.2, True, DARK_BLUE)
    r = p.add_run("A: " + a)
    set_run_font(r, 10.2, False, INK)

add_quick_reference(doc)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.runs:
                    para_format(p, after=0, line=1.12)

doc.core_properties.title = "SVdP Voucher System How-To Guide for Cashiers"
doc.core_properties.subject = "Cashier guide for ordinary voucher lookup, redemption, emergency voucher, coat, furniture, and household goods workflows"
doc.core_properties.author = "Society of St. Vincent de Paul - Fort Wayne"
doc.save(OUT)
print(OUT.resolve())
