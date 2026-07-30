from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/docs/SVdP Voucher System How-To Guide - Website Admin.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 88, 88)
INK = RGBColor(0, 0, 0)
LIGHT_BLUE = "E8EEF5"
PALE = "F4F6F9"
WARNING = "FFF2CC"


def set_run_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para_format(p, before=0, after=6, line=1.2, left=None, first=None, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if left is not None:
        p.paragraph_format.left_indent = Inches(left)
    if first is not None:
        p.paragraph_format.first_line_indent = Inches(first)
    p.paragraph_format.keep_together = keep
    p.paragraph_format.keep_with_next = keep


def add_para(doc, text="", size=10.5, bold=False, color=INK, before=0, after=6, align=None):
    p = doc.add_paragraph()
    para_format(p, before, after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size, bold, color)
    return p


def add_title_block(doc):
    add_para(doc, "FOR WEBSITE ADMIN USE", 9, True, MUTED, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "THE SOCIETY OF ST. VINCENT DE PAUL", 14, True, DARK_BLUE, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Neighbors helping Neighbors", 10.5, False, MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Voucher System", 24, True, INK, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "How-To Guide for Website Admins", 15, False, BLUE, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 7
    elif level == 2:
        size, color, before, after = 13, BLUE, 10, 5
    else:
        size, color, before, after = 11.5, DARK_BLUE, 8, 3
    para_format(p, before, after, line=1.15, keep=True)
    r = p.add_run(text)
    set_run_font(r, size, True, color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    indent = 0.44 + (level * 0.24)
    para_format(p, 0, 3, line=1.18, left=indent, first=-0.22)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_step(doc, number, text):
    p = doc.add_paragraph()
    para_format(p, 0, 3, line=1.18, left=0.05)
    label = p.add_run(f"Step {number}: ")
    set_run_font(label, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_steps(doc, steps, start=1):
    for i, step in enumerate(steps, start):
        add_step(doc, i, step)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=9.45, bold=False, color=INK):
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.12)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for cell, head in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        set_cell_text(cell, head, 9.4, True, DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_data):
            set_cell_text(cell, text)
    return table


def add_note(doc, label, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.18)
    r = p.add_run(label + ": ")
    set_run_font(r, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_code_line(doc, text):
    p = doc.add_paragraph()
    para_format(p, before=2, after=7, line=1.12, left=0.25)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK, name="Courier New")


def add_quick_reference(doc):
    doc.add_page_break()
    add_heading(doc, "Quick Reference Card", 1)
    add_para(doc, "Use this as a setup reminder after you have read the full guide.", 10.5, False, MUTED, after=8)
    rows = [
        ("Create organization", "SVdP Vouchers -> Conferences -> Add New Conference or Partner -> Add Organization."),
        ("Copy slug", "Existing Organizations -> Slug column. Use the exact slug shown there."),
        ("Approve voucher types", "Existing Organizations -> Edit Types -> select approved types -> Save Types."),
        ("Add notification email", "Existing Organizations -> Notification Email -> Update."),
        ("Add billing email", "Accounting -> Organization Billing Mappings -> Billing Email -> Save."),
        ("Create page", "WordPress Pages -> Add New Page -> add shortcode -> Publish."),
        ("Shortcode format", "[svdp_voucher_request conference=\"slug-here\"]"),
        ("Test page", "Open the published page and confirm the organization name appears on the form."),
    ]
    add_table(doc, ["TASK", "WHERE / WHAT TO DO"], rows, [2450, 6910])
    add_heading(doc, "Important Rules", 2)
    add_bullet(doc, "The shortcode slug must match the organization slug exactly.")
    add_bullet(doc, "If the slug is wrong, the page will show Conference not found.")
    add_bullet(doc, "Voucher types control what the organization can request.")
    add_bullet(doc, "Billing Email is on the Accounting tab, not the Conferences tab.")
    add_bullet(doc, "Deleting an organization makes it unavailable for new vouchers. It does not erase old vouchers.")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.45)
section.footer_distance = Inches(0.38)

for style_name in ["Normal", "List Bullet", "List Number"]:
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_run = header.add_run("SVdP Voucher System - Website Admin Guide")
set_run_font(header_run, 8.5, False, MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("(c) 2026 Society of St. Vincent de Paul - Fort Wayne. All rights reserved.")
set_run_font(footer_run, 8.5, False, MUTED)

add_title_block(doc)
add_heading(doc, "What This Guide Covers", 1)
add_para(doc, "This guide explains how a Website Admin sets up a new Conference or Partner request page and connects that page to the Voucher System.", 10.5)
add_bullet(doc, "Create the Conference or Partner in the Voucher System.")
add_bullet(doc, "Confirm the slug that will be used on the WordPress page.")
add_bullet(doc, "Approve the voucher types the organization can request.")
add_bullet(doc, "Add the billing email used for statements and bookkeeping.")
add_bullet(doc, "Create a WordPress page with the correct shortcode.")
add_note(doc, "Important", "Do the Voucher System setup first. Then create the WordPress page. The page needs the organization slug before it can pre-fill the form correctly.", WARNING)

add_heading(doc, "How the Pieces Fit Together", 1)
add_table(doc, ["PIECE", "WHAT IT DOES"], [
    ("Organization record", "Stores the Conference or Partner name, slug, eligibility window, item limit, voucher types, and notification email."),
    ("Slug", "The short name used by the shortcode to find the correct organization."),
    ("Voucher types", "Controls whether the organization can request clothing, furniture, household goods, or only selected types."),
    ("Billing Email", "The email used for statement delivery and bookkeeping work."),
    ("WordPress page", "The public request page shared with the Conference or Partner."),
    ("Shortcode", "The small line placed on the page to display the request form."),
], [2450, 6910])

doc.add_page_break()
add_heading(doc, "Create a New Conference or Partner", 1)
add_para(doc, "Start on the Conferences tab. This creates the organization record the form will use.", 10.5)
add_steps(doc, [
    "Sign in to WordPress with Website Admin access.",
    "Open the WordPress Dashboard.",
    "Go to SVdP Vouchers.",
    "Open the Conferences tab.",
    "Find Add New Conference or Partner.",
    "Choose Organization Type: Conference or Partner.",
    "Enter the Organization Name.",
    "Review the Slug. Leave it blank to let the system create one, or enter the slug you want to use.",
    "Set Eligibility Window if this organization needs a different number of days between requests.",
    "Set Items Per Person if this organization needs a different item limit.",
    "Click Add Organization.",
])
add_bullet(doc, "Conferences are St. Vincent de Paul organizations.")
add_bullet(doc, "Partners are external organizations that may be billed.")
add_bullet(doc, "The system will reload the list after the organization is created.")

add_heading(doc, "Check the Slug", 2)
add_steps(doc, [
    "Stay on the Conferences tab.",
    "Find the organization under Existing Organizations.",
    "Look at the Slug column.",
    "Use that exact slug in the shortcode on the WordPress page.",
])
add_note(doc, "Example", "If the slug is st-mary-fort-wayne, the shortcode must use st-mary-fort-wayne.", PALE)
add_bullet(doc, "Do not use the organization name in the shortcode.")
add_bullet(doc, "Do not add extra spaces inside the slug.")

add_heading(doc, "Update Basic Organization Details", 2)
add_steps(doc, [
    "Find the organization under Existing Organizations.",
    "Update Eligibility or Items/Person if needed.",
    "Enter the Notification Email if voucher request notices should go to the organization.",
    "Click Update.",
])
add_bullet(doc, "Notification Email is for voucher request notifications.")
add_bullet(doc, "Billing Email is handled separately on the Accounting tab.")

doc.add_page_break()
add_heading(doc, "Approve Voucher Types", 1)
add_para(doc, "Voucher types decide what the organization is allowed to request on its form.", 10.5)
add_heading(doc, "Set Allowed Types for One Organization", 2)
add_steps(doc, [
    "Open the Conferences tab.",
    "Find the organization under Existing Organizations.",
    "Click Edit Types.",
    "Check each voucher type this organization is approved to request.",
    "Uncheck any voucher type this organization should not request.",
    "Click Save Types.",
    "Wait for the success message and page reload.",
])
add_bullet(doc, "Available choices can include Clothing, Furniture, and Household Goods.")
add_bullet(doc, "At least one voucher type must be selected.")
add_bullet(doc, "The selected types are shown under the Edit Types button after the page reloads.")
add_note(doc, "Use care", "Only approve voucher types that the Conference or Partner is allowed to offer. These choices affect what users can request from that organization's page.", WARNING)

add_heading(doc, "If a Voucher Type Is Missing", 2)
add_para(doc, "The Edit Types box only shows voucher types that are available system-wide.", 10.5)
add_steps(doc, [
    "Open the Settings tab.",
    "Find Voucher Types Management.",
    "Review Available Voucher Types.",
    "Confirm the needed type is listed.",
    "Review whether delivery can be selected for that type if delivery applies.",
    "Click Save All Settings if changes are made.",
])
add_bullet(doc, "Release C supports clothing, furniture, and household_goods.")
add_bullet(doc, "After a system-wide voucher type is available, return to Conferences and approve it for the organization.")

add_heading(doc, "What the Request Form Does With Voucher Types", 2)
add_bullet(doc, "The form reads the organization's approved voucher types.")
add_bullet(doc, "The user can only choose from the approved types.")
add_bullet(doc, "If the organization is approved for one type, the form is limited to that type.")
add_bullet(doc, "If the organization is approved for several types, the form lets the user choose among those types.")

doc.add_page_break()
add_heading(doc, "Add the Billing Email", 1)
add_para(doc, "Billing Email is used for statements and bookkeeping. It is not entered on the Conferences tab.", 10.5)
add_steps(doc, [
    "Open the Accounting tab.",
    "Find Organization Billing Mappings.",
    "Find the Conference or Partner.",
    "Enter the Billing Email.",
    "Enter the QuickBooks Customer Name if bookkeeping uses QuickBooks for this organization.",
    "Click Save on that organization row.",
])
add_bullet(doc, "Use the billing address the organization wants statements sent to.")
add_bullet(doc, "If Billing Email is missing, statement sending may fall back to the notification email when available.")
add_bullet(doc, "The QuickBooks Customer Name should match the customer name used in QuickBooks.")
add_note(doc, "Remember", "Notification Email and Billing Email serve different purposes. Notification Email is for voucher request notices. Billing Email is for statements and bookkeeping.", PALE)

add_heading(doc, "Create the WordPress Page", 1)
add_para(doc, "Each Conference or Partner can have its own request page. The shortcode connects the page to that organization.", 10.5)
add_steps(doc, [
    "Open the WordPress Dashboard.",
    "Go to Pages.",
    "Click Add New Page.",
    "Enter a clear page title, such as Clothing Voucher Request - St Mary.",
    "Add a shortcode block or plain paragraph block.",
    "Enter the voucher request shortcode using the organization's slug.",
])
add_code_line(doc, '[svdp_voucher_request conference="slug-here"]')
add_steps(doc, [
    "Replace slug-here with the exact slug from the Conferences tab.",
    "Publish the page.",
    "Copy the published page URL.",
    "Share the URL with the Conference or Partner.",
], start=7)

add_heading(doc, "Shortcode Examples", 2)
add_table(doc, ["ORGANIZATION SLUG", "SHORTCODE"], [
    ("st-mary-fort-wayne", '[svdp_voucher_request conference="st-mary-fort-wayne"]'),
    ("community-partner", '[svdp_voucher_request conference="community-partner"]'),
    ("st-joseph", '[svdp_voucher_request conference="st-joseph"]'),
], [2650, 6710])

doc.add_page_break()
add_heading(doc, "Test the Page", 1)
add_para(doc, "Always test the page before sharing it.", 10.5)
add_steps(doc, [
    "Open the published page in a browser.",
    "Confirm the voucher request form appears.",
    "Go to the Requestor step.",
    "Confirm the Organization line shows the correct Conference or Partner name.",
    "Confirm the organization dropdown is not shown.",
    "Confirm the Assistance Needed choices match the approved voucher types.",
    "Submit only if you are intentionally creating a real test request. Otherwise, stop after reviewing the form.",
])
add_note(doc, "If the page shows an error", "Check the shortcode slug. If the slug does not match an organization in the system, the page will show Conference not found.", WARNING)

add_heading(doc, "Deactivate an Organization", 1)
add_para(doc, "Use this when a Conference or Partner should no longer appear for new voucher requests.", 10.5)
add_steps(doc, [
    "Open the Conferences tab.",
    "Find the organization under Existing Organizations.",
    "Click the delete button for that organization.",
    "Confirm the message.",
])
add_bullet(doc, "The organization is made unavailable for new vouchers.")
add_bullet(doc, "Existing vouchers are not erased.")
add_bullet(doc, "The Store row is system managed and cannot be removed from this screen.")

add_heading(doc, "Common Questions", 1)
qa = [
    ("What is a slug?", "It is the short page-friendly name used by the shortcode to find the right organization."),
    ("Can I change the slug later?", "The system can store a changed slug, but the WordPress page shortcode must be updated to match it."),
    ("What happens if the slug is wrong?", "The request page will show Conference not found."),
    ("Where do I add the billing email?", "Use the Accounting tab, under Organization Billing Mappings."),
    ("Where do I add the notification email?", "Use the Conferences tab, under Existing Organizations."),
    ("How do I control what the organization can request?", "Use Edit Types on the Conferences tab and select the approved voucher types."),
    ("Do I need one page per organization?", "Yes, when each Conference or Partner should have a direct link that opens with its name already selected."),
]
for question, answer in qa:
    add_para(doc, question, 10.3, True, DARK_BLUE, before=4, after=2)
    add_para(doc, answer, 10.2, False, INK, after=4)

add_quick_reference(doc)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
