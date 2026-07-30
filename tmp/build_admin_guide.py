from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/docs/SVdP Voucher System How-To Guide - Administrators and Managers.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 88, 88)
INK = RGBColor(0, 0, 0)
LIGHT_BLUE = "E8EEF5"
PALE = "F4F6F9"
WARNING = "FFF2CC"


def set_run_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para_format(p, before=0, after=6, line=1.2, left=None, first=None, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if left is not None:
        p.paragraph_format.left_indent = Inches(left)
    if first is not None:
        p.paragraph_format.first_line_indent = Inches(first)
    p.paragraph_format.keep_together = keep
    p.paragraph_format.keep_with_next = keep


def add_para(doc, text="", size=10.5, bold=False, color=INK, before=0, after=6, align=None):
    p = doc.add_paragraph()
    para_format(p, before, after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size, bold, color)
    return p


def add_title_block(doc):
    add_para(doc, "FOR ADMINISTRATOR & MANAGER USE ONLY", 9, True, MUTED, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "THE SOCIETY OF ST. VINCENT DE PAUL", 14, True, DARK_BLUE, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Neighbors helping Neighbors", 10.5, False, MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Voucher System", 24, True, INK, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "How-To Guide for Administrators & Managers", 15, False, BLUE, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 7
    elif level == 2:
        size, color, before, after = 13, BLUE, 10, 5
    else:
        size, color, before, after = 11.5, DARK_BLUE, 8, 3
    para_format(p, before, after, line=1.15, keep=True)
    r = p.add_run(text)
    set_run_font(r, size, True, color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    indent = 0.44 + (level * 0.24)
    para_format(p, 0, 3, line=1.18, left=indent, first=-0.22)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_step(doc, number, text):
    p = doc.add_paragraph()
    para_format(p, 0, 3, line=1.18, left=0.05)
    label = p.add_run(f"Step {number}: ")
    set_run_font(label, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, size=9.45, bold=False, color=INK):
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.12)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for cell, head in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        set_cell_text(cell, head, 9.4, True, DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_data):
            set_cell_text(cell, text)
    return table


def add_note(doc, label, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    p = cell.paragraphs[0]
    para_format(p, after=0, line=1.18)
    r = p.add_run(label + ": ")
    set_run_font(r, 10.2, True, DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, 10.2, False, INK)


def add_quick_reference(doc):
    doc.add_page_break()
    add_heading(doc, "Quick Reference Card", 1)
    add_para(doc, "Use this as a station-side reminder after you have read the full guide.", 10.5, False, MUTED, after=8)
    rows = [
        ("Add manager", "Managers tab -> enter name -> optional 4-character code -> Add Manager -> save the code right away."),
        ("Regenerate code", "Managers tab -> Regenerate Code -> save the new code. The old code stops working immediately."),
        ("Deactivate manager", "Managers tab -> Deactivate. Their code no longer works."),
        ("Override reasons", "Override Reasons tab -> add, edit, delete, or drag to reorder the dropdown."),
        ("Furniture catalog", "Furniture Catalog tab -> manage categories and catalog items. Archive instead of deleting."),
        ("Household goods", "Household Goods tab -> manage browse groups, categories, limits, and configuration audit."),
        ("Voucher edit", "Cashier Station -> open voucher -> Correct Voucher -> enter changes -> manager name, code, and reason -> Submit Correction."),
        ("Emergency override", "Cashier Station -> duplicate found -> manager enters name, code, reason -> Validate & Create."),
        ("Audit review", "Voucher Correction Audit tab -> filter by voucher, neighbor, field, manager, actor, reason, or date."),
    ]
    add_table(doc, ["ACTION", "WHERE / WHAT TO DO"], rows, [2400, 6960])
    add_heading(doc, "Important Rules", 2)
    add_bullet(doc, "Manager codes are 4 characters using A-Z and 2-9.")
    add_bullet(doc, "Codes are shown only when created or regenerated. Save them securely.")
    add_bullet(doc, "Regenerating a code replaces the old code immediately.")
    add_bullet(doc, "Archive catalog rows and cancellation reasons instead of deleting them, so old vouchers and reports still show what happened at the time.")
    add_bullet(doc, "Use the audit log when you need to understand who changed what, when, and why.")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.45)
section.footer_distance = Inches(0.38)

for style_name in ["Normal", "List Bullet", "List Number"]:
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.text = "SVdP Voucher System - Administrator & Manager Guide"
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(header.runs[0], 8.5, False, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("(c) 2026 Society of St. Vincent de Paul - Fort Wayne. All rights reserved.")
set_run_font(r, 8.5, False, MUTED)

add_title_block(doc)

add_heading(doc, "What You Need to Know", 1)
add_para(doc, "This guide is for people who manage the Voucher System configuration and approve special cashier actions. It explains the admin tabs, the catalog tools, manager codes, reason lists, audit review, and cashier-station actions that need manager approval or careful supervisor review.", 10.5)
add_bullet(doc, "Most admin work happens in WordPress under the SVdP Vouchers admin page.")
add_bullet(doc, "Most manager-code work happens in the Cashier Station when the system asks for approval.")
add_bullet(doc, "Changes to catalogs and settings affect future requests. Existing voucher snapshots keep their historical information.")
add_note(doc, "Important", "Manager codes should be treated like keys. Give them only to people who are allowed to approve exceptions or voucher corrections.", WARNING)

add_heading(doc, "Opening the Admin Area", 1)
add_steps(doc, [
    "Sign in to WordPress with an account that has permission to manage the Voucher System.",
    "Open the WordPress Dashboard.",
    "Go to the SVdP Vouchers admin page.",
    "Use the tabs across the top to choose the area you need.",
])

add_heading(doc, "Admin Tabs at a Glance", 1)
add_table(doc, ["TAB", "WHAT IT IS FOR"], [
    ("Analytics", "Reporting, voucher activity, exports, and operational review."),
    ("Conferences", "Conference and Partner organization records, access, and request form details."),
    ("Furniture Catalog", "Furniture categories and furniture catalog items."),
    ("Furniture Reasons", "Preset cancellation reasons used when furniture items cannot be fulfilled."),
    ("Household Goods", "Browse groups, Household Goods categories, request limits, and configuration audit."),
    ("Invoices / Statements / Accounting", "Billing and accounting review for fulfilled furniture and household goods work."),
    ("Managers", "Manager names, active status, and manager codes."),
    ("Override Reasons", "Reasons shown when manager approval is required."),
    ("Voucher Correction Audit", "Read-only history of manager-authorized voucher corrections."),
    ("Settings", "Item values, voucher types, delivery options, form text, database counts, and system information."),
], [2450, 6910])

doc.add_page_break()
add_heading(doc, "Managing the Furniture Catalog", 1)
add_para(doc, "The Furniture Catalog controls what Conferences and Partners can request and what the store team sees during furniture fulfillment.", 10.5)
add_heading(doc, "Furniture Categories", 2)
add_steps(doc, [
    "Open the Furniture Catalog tab.",
    "In Furniture Categories, enter the category name.",
    "Enter a sort order if you want the category to appear in a specific order.",
    "Click Add Category.",
])
add_bullet(doc, "Use Edit to change a category name or sort order.")
add_bullet(doc, "Use Archive to hide a category from future use.")
add_bullet(doc, "A category cannot be archived while active items remain in it.")
add_heading(doc, "Furniture Catalog Items", 2)
add_steps(doc, [
    "Open the Furniture Catalog tab.",
    "In Furniture Catalog, enter the item name.",
    "Choose the category.",
    "Choose the pricing type: Range or Fixed.",
    "Enter the price fields shown for that pricing type.",
    "Choose whether public range pricing should show as Up to the maximum amount.",
    "Choose the Conference Coverage Type: Percent or Fixed Dollar Amount.",
    "Enter the coverage value and sort order.",
    "Click Add Catalog Item.",
])
add_bullet(doc, "Use range pricing for used furniture or items where final price can vary.")
add_bullet(doc, "Use fixed pricing for handmade furniture, mattresses, frames, or items with one set price.")
add_bullet(doc, "Furniture catalog items can be archived, but they are not deleted from the system.")
add_bullet(doc, "This protects old voucher records. If a neighbor requested a Sofa six months ago, that old voucher still needs to show Sofa even if Sofa is no longer offered for new requests.")
add_bullet(doc, "Use Archive when an item should disappear from future request forms. Use Restore if it should be offered again.")

add_heading(doc, "Furniture Cancellation Reasons", 1)
add_para(doc, "Furniture Cancellation Reasons are the preset choices used when a requested furniture item cannot be provided.", 10.5)
add_steps(doc, [
    "Open the Furniture Reasons tab.",
    "Enter the Reason Text.",
    "Enter a Display Order if you want to control where it appears.",
    "Click Add Cancellation Reason.",
])
add_bullet(doc, "Use Edit to change the reason text or display order.")
add_bullet(doc, "Use Archive when a reason should no longer be used.")
add_bullet(doc, "Archive old reasons instead of deleting them so past furniture records stay readable.")
add_bullet(doc, "This keeps old records clear because the reason selected at the time can still be understood later.")

add_heading(doc, "Managing the Household Goods Catalog", 1)
add_para(doc, "Household Goods uses browse groups and requestable categories. Browse groups control the filter pills. Categories are the actual household goods needs that can be requested.", 10.5)
add_heading(doc, "Browse Groups", 2)
add_steps(doc, [
    "Open the Household Goods tab.",
    "In Household Goods Browse Groups, enter the group name.",
    "Enter a sort order.",
    "Click Add Browse Group.",
])
add_bullet(doc, "Use Edit to change the name or sort order.")
add_bullet(doc, "Use Archive to hide a browse group from future requests.")
add_bullet(doc, "Browse groups can be archived when they should no longer appear for new requests.")
add_bullet(doc, "They are kept in the system so old Household Goods requests can still be read accurately later.")
add_heading(doc, "Household Goods Categories", 2)
add_steps(doc, [
    "Open the Household Goods tab.",
    "In Household Goods Catalog, enter the category name.",
    "Choose the browse group.",
    "Choose pricing type: Fixed or Range.",
    "Enter the retail price fields shown for that pricing type.",
    "Choose whether public range pricing should show as Up to the maximum amount.",
    "Choose Organization Coverage Type: Percent or Fixed Dollar Amount.",
    "Enter Organization Coverage, Quantity Maximum, Sort Order, and Cashier Guidance if needed.",
    "Click Add Household Goods Category.",
])
add_bullet(doc, "Quantity Maximum limits how many of that category can be requested. Use 0 for no category-level limit.")
add_bullet(doc, "Cashier Guidance is for managed staff guidance only. It is not shown to Vincentians.")
add_bullet(doc, "Household Goods categories can be archived, but they are not deleted from the system.")
add_bullet(doc, "This protects old voucher records. If a neighbor requested Bath Towels before the category was archived, that old voucher still needs to show Bath Towels.")
add_bullet(doc, "Use Archive when a category should disappear from future request forms. Use Restore if it should be offered again.")
add_heading(doc, "Household Goods Limits", 2)
add_steps(doc, [
    "Open the Household Goods tab.",
    "Find Household Goods Limits.",
    "Set Maximum Selected Categories.",
    "Set Maximum Total Requested Quantity.",
    "Use 0 when no limit should apply.",
    "Click Save Limits.",
])
add_heading(doc, "Configuration Audit", 2)
add_bullet(doc, "The Household Goods Configuration Audit is read-only.")
add_bullet(doc, "It shows what changed, who changed it, before and after values, and a summary.")

add_heading(doc, "Managing Managers and Codes", 1)
add_para(doc, "Managers approve emergency voucher overrides and voucher corrections. Each active manager has one unique 4-character code.", 10.5)
add_heading(doc, "Add a New Manager", 2)
add_steps(doc, [
    "Open the Managers tab.",
    "Enter the Manager Name.",
    "Optional: enter a 4-character code using A-Z and 2-9.",
    "If you leave the code blank, the system generates one.",
    "Click Add Manager.",
    "When the code appears, save it securely right away.",
])
add_note(doc, "Important", "The code is shown only when the manager is created or when the code is regenerated.", WARNING)
add_heading(doc, "Regenerate a Manager Code", 2)
add_steps(doc, [
    "Open the Managers tab.",
    "Find the active manager.",
    "Click Regenerate Code.",
    "Confirm that you want to generate a new code.",
    "Save the new code securely.",
])
add_bullet(doc, "The old code stops working immediately.")
add_bullet(doc, "Regenerating a code also clears failed attempts and lockout for that manager.")
add_heading(doc, "Deactivate a Manager", 2)
add_steps(doc, [
    "Open the Managers tab.",
    "Find the active manager.",
    "Click Deactivate.",
    "Confirm the deactivation.",
])
add_bullet(doc, "A deactivated manager code no longer works.")
add_bullet(doc, "Use this when someone should no longer approve exceptions.")

add_heading(doc, "Override Reasons", 1)
add_para(doc, "Override Reasons appear in the dropdown when manager approval is required for an emergency voucher override or voucher correction.", 10.5)
add_steps(doc, [
    "Open the Override Reasons tab.",
    "To add a reason, enter Reason Text and click Add Reason.",
    "To edit a reason, click Edit, update the text, and click Save.",
    "To remove a reason from the list, click Delete.",
    "To reorder reasons, drag the rows into the order you want.",
])
add_bullet(doc, "Use simple reasons that explain the decision clearly.")
add_bullet(doc, "The selected reason is saved with the approval attempt or correction record.")

add_heading(doc, "Voucher Correction Audit", 1)
add_para(doc, "The Voucher Correction Audit is a read-only history of manager-authorized voucher corrections.", 10.5)
add_heading(doc, "How to Filter the Audit", 2)
add_steps(doc, [
    "Open the Voucher Correction Audit tab.",
    "Use any filter you need: Voucher ID, Neighbor, Field, Manager, Actor, Reason, Date From, Date To, or Per Page.",
    "Click Filter.",
    "Use Clear Filters to return to the full list.",
])
add_heading(doc, "How to Read an Audit Row", 2)
add_table(doc, ["FIELD", "WHAT IT TELLS YOU"], [
    ("Summary", "Plain-language description of what changed."),
    ("Voucher", "The voucher number and neighbor connected to the change."),
    ("Field", "Which voucher field changed."),
    ("Manager", "The manager name recorded with the approval."),
    ("Actor", "The logged-in user who submitted the change."),
    ("Reason", "The selected reason at the time of approval."),
    ("Conference", "The Conference or Partner tied to the voucher."),
    ("Timestamp", "When the correction was recorded."),
], [2200, 7160])

add_heading(doc, "Settings Page", 1)
add_para(doc, "The Settings tab controls system-wide text and values. Be careful here: settings can change what users see and how redemption values are calculated.", 10.5)
add_table(doc, ["SECTION", "WHAT IT CONTROLS"], [
    ("Item Value Configuration", "Adult Item Value and Child Item Value used when clothing vouchers are redeemed."),
    ("Example Calculation", "Shows how adult and child item values calculate an estimated redemption total."),
    ("Voucher Types Management", "Available root voucher types, delivery availability, and Assistance Needed descriptions."),
    ("Available Voucher Types", "Comma-separated root types. Current supported types are clothing, furniture, and household_goods."),
    ("Delivery Available", "Whether delivery can be selected for each voucher type."),
    ("Assistance Step Description", "The short description shown on the public request form for each voucher type."),
    ("Form Boilerplate Text", "Store Hours and Redemption Instructions shown on request forms."),
    ("Database Info", "Current voucher and active Conference counts."),
    ("System Information", "Plugin version, WordPress version, and PHP version."),
], [2600, 6760])
add_steps(doc, [
    "Open the Settings tab.",
    "Make the needed changes.",
    "Click Save All Settings.",
    "Wait for the settings saved message.",
])
add_note(doc, "Use care", "Changing voucher types, delivery settings, or request-form language can affect future Conference and Partner requests.", WARNING)

doc.add_page_break()
add_heading(doc, "Manager Approval in the Cashier Station", 1)
add_para(doc, "Manager codes are used in the Cashier Station when the system asks for approval. The manager enters their name, 4-character code, and a reason.", 10.5)
add_heading(doc, "Override an Emergency Voucher Duplicate", 2)
add_steps(doc, [
    "A cashier fills out an Emergency Voucher.",
    "If an exact duplicate is found, the system opens Duplicate Found: Manager Approval Required.",
    "Review the neighbor, issuing organization, created date, and next eligible date.",
    "Enter the manager code.",
    "Enter the manager name exactly as it appears in the Managers list.",
    "Choose the reason.",
    "Click Validate & Create.",
])
add_bullet(doc, "If approved, the emergency clothing voucher is created.")
add_bullet(doc, "If cancelled, the duplicate denial can be saved for tracking.")
add_heading(doc, "Edit a Voucher", 2)
add_para(doc, "Voucher correction is used when voucher details need to be changed after creation. It requires manager approval and creates an audit row.", 10.5)
add_steps(doc, [
    "In the Cashier Station, search for and open the voucher.",
    "Open the Correct Voucher section.",
    "Update only the fields that need correction.",
    "Choose the correction reason.",
    "Enter the manager name.",
    "Enter the manager code.",
    "Click Submit Correction.",
    "Confirm the success message and review the voucher details.",
])
add_bullet(doc, "Editable fields can include household counts, date of birth, voucher created date, status, and delivery address fields when present.")
add_bullet(doc, "The audit records the manager, actor, reason, field changed, before value, after value, and timestamp.")
add_heading(doc, "Cancel a Furniture Item", 2)
add_para(doc, "Furniture item cancellation uses the Furniture Cancellation Reasons list. In ordinary flow, the manager code is not entered on the cancellation form, but this should still be handled with supervisor judgment.", 10.5)
add_steps(doc, [
    "Open the furniture voucher in the Cashier Station.",
    "Find the requested furniture item.",
    "Click Cancel Item.",
    "Choose the Cancellation Reason.",
    "Add notes if helpful.",
    "Click Confirm Cancellation.",
])
add_bullet(doc, "Use cancellation only when the item cannot be provided.")
add_bullet(doc, "The selected cancellation reason comes from the Furniture Reasons admin tab.")
add_heading(doc, "Mark Household Goods Unavailable", 2)
add_para(doc, "Household Goods uses unavailable quantities instead of the furniture Cancel Item button.", 10.5)
add_steps(doc, [
    "Open the Household Goods voucher in the Cashier Station.",
    "Find the requested line.",
    "Enter the unavailable quantity.",
    "Choose the unavailable reason when the unavailable quantity is greater than zero.",
    "Use Save Progress if the voucher is not finished.",
    "Finalize only when all requested units are resolved.",
])
add_bullet(doc, "Fulfilled plus unavailable quantity cannot exceed the requested quantity.")
add_bullet(doc, "Once finalized, the voucher cannot be edited through the ordinary cashier workflow.")

add_heading(doc, "Common Questions", 1)
qa = [
    ("What if a manager forgot their code?", "Regenerate the code from the Managers tab and give them the new code securely."),
    ("Can I see an old manager code?", "No. Codes are shown only when created or regenerated."),
    ("What happens if a manager is deactivated?", "Their code stops working. Their historical approvals remain in the records."),
    ("Should I delete or archive catalog records?", "Archive them. Catalog records are connected to old vouchers, and old vouchers need to keep showing what was requested at the time."),
    ("Where do emergency override reasons come from?", "They come from the Override Reasons tab."),
    ("Where do furniture cancellation reasons come from?", "They come from the Furniture Reasons tab."),
    ("Where do I review voucher edits?", "Use the Voucher Correction Audit tab."),
    ("What if the Cashier Station asks for approval and the code fails?", "Check the manager name, code, and reason. After repeated failed attempts, the manager may be temporarily locked out."),
]
for q, a in qa:
    p = doc.add_paragraph()
    para_format(p, before=2, after=2, line=1.16)
    r = p.add_run("Q: " + q + " ")
    set_run_font(r, 10.2, True, DARK_BLUE)
    r = p.add_run("A: " + a)
    set_run_font(r, 10.2, False, INK)

add_quick_reference(doc)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.runs:
                    para_format(p, after=0, line=1.12)

doc.core_properties.title = "SVdP Voucher System How-To Guide for Administrators and Managers"
doc.core_properties.subject = "Administrator and manager guide for voucher system configuration, manager codes, catalogs, reasons, settings, audit, and cashier approvals"
doc.core_properties.author = "Society of St. Vincent de Paul - Fort Wayne"
doc.save(OUT)
print(OUT.resolve())
